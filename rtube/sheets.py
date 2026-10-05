"""Пари «ID товару — посилання на відео» з файлу: xlsx / xlsm, csv, txt.

Колонки шукаємо самі, а не за жорсткою схемою: у кожному рядку — посилання на
YouTube (зокрема гіперпосилання під текстом «відео») і ID товару (число з
5–12 цифр). Розбирає рядок та сама tools.extract_id_pairs, що й вставлений
список, тож правила однакові.

Якщо в шапці є стовпець з «ID»/«ІД» у назві, ID беремо саме з нього: у
таблиці може бути ще якесь число (код продавця, штрихкод), і «перше число
в рядку» тоді вгадало б не те.
"""

import csv
import io
import os
import re
from dataclasses import dataclass, field

from . import tools

SUPPORTED = (".xlsx", ".xlsm", ".csv", ".txt")
_HEADER_ID = re.compile(r"(^|[^a-zа-яіїєґ])(id|ід|ид)([^a-zа-яіїєґ]|$)", re.I)
_ID_CELL = re.compile(r"^\s*(\d{5,12})\s*$")


@dataclass
class SheetResult:
    pairs: list = field(default_factory=list)   # [(ID або None, посилання), …]
    rows: int = 0                               # непорожніх рядків прочитано
    skipped: int = 0                            # рядків без посилання на ролик
    sheet: str = ""

    @property
    def with_ids(self):
        return sum(1 for pid, _ in self.pairs if pid)


def _cell_text(value):
    """Число з Excel (590312170.0) → «590312170»; решта — як є."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _find_id_column(header):
    for i, text in enumerate(header):
        if text and not tools.extract_video_urls(text) and _HEADER_ID.search(text):
            return i
    return None


def pairs_from_rows(rows):
    """rows — [[текст клітинки, …], …] (гіперпосилання вже додано в текст).
    Повертає SheetResult."""
    result = SheetResult()
    id_col = None
    for index, cells in enumerate(rows):
        if not any(cells):
            continue
        result.rows += 1
        line = "\t".join(cells)
        if not tools.extract_video_urls(line):
            # Шапка (рядок без посилань серед перших) підказує колонку з ID —
            # і пропущеним рядком тоді не рахується.
            if id_col is None and index < 5:
                id_col = _find_id_column(cells)
                if id_col is not None:
                    continue
            result.skipped += 1
            continue
        found = tools.extract_id_pairs(line)
        if id_col is not None and id_col < len(cells):
            match = _ID_CELL.match(cells[id_col])
            product_id = match.group(1) if match else None
            found = [(product_id, url) for _, url in found]
        result.pairs.extend(found)
    return result


def _xlsx_rows(path):
    from openpyxl import load_workbook
    # Не read_only: лише так видно гіперпосилання клітинок.
    wb = load_workbook(path, data_only=True)
    try:
        sheets = [wb.active] + [ws for ws in wb.worksheets if ws is not wb.active]
        for ws in sheets:
            rows = []
            for row in ws.iter_rows():
                cells = []
                for cell in row:
                    text = _cell_text(cell.value)
                    link = getattr(cell.hyperlink, "target", None) if cell.hyperlink else None
                    if link and link not in text:
                        text = f"{text} {link}".strip()
                    cells.append(text)
                rows.append(cells)
            yield ws.title, rows
    finally:
        wb.close()


def _text_rows(path):
    with open(path, "rb") as f:
        raw = f.read()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = raw.decode("utf-8", "replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        rows = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    except csv.Error:
        rows = [[line] for line in text.splitlines()]
    yield os.path.basename(path), rows


def read_pairs(path):
    """Пари з файлу. Для книги Excel — перший аркуш, де є посилання на ролики
    (спершу активний). Кидає ValueError з поясненням, якщо нічого не знайдено."""
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED:
        raise ValueError(f"Непідтримуваний файл {ext or '(без розширення)'} — "
                         "потрібен .xlsx, .csv або .txt")
    sources = _xlsx_rows(path) if ext in (".xlsx", ".xlsm") else _text_rows(path)
    for name, rows in sources:
        result = pairs_from_rows(rows)
        if result.pairs:
            result.sheet = name
            return result
    raise ValueError("У файлі не знайдено жодного посилання на відео YouTube")
