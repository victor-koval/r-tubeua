"""Звіт по завантаженнях у xlsx: що вийшло з пакета — до заливання на FTP.

Рядок на кожне завдання зі списку: ID товару, посилання, ім'я файлу, статус.
Тривалість — текстом «3,27» (хвилин,секунд), як її копіює клік у програмі:
числом Excel показав би «3,2» замість «3,20».

Товари з тим самим роликом ділять один файл (названий за першим ID): у
«Файлі» в них той самий файл, у «Деталях» — з ким саме, а рядки однієї
групи зафарбовано одним кольором.

«Роздільна здатність» — розмір кадру готового файлу. Нижчу за вимоги Rozetka
(квадратне — від 640×640, горизонтальне й вертикальне — від 720p) підсвічено
червоним, а внизу — скільки таких: на пакеті з сотень роликів їх інакше не
видно.
"""

import os
import time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .uikit import format_min_sec, whole_seconds

HEADERS = ("№", "ID товару", "Посилання", "Назва", "Файл", "Статус", "Деталі",
           "Тривалість", "Роздільна здатність", "Розмір, МБ", "Відео на Rozetka")
RESOLUTION_COL = HEADERS.index("Роздільна здатність") + 1
MAX_WIDTH = 60

ALREADY_NOTE = "Уже є в теці"     # початок downloader.ALREADY_NOTE
SAME_VIDEO = "Те саме відео"      # початок downloader.same_video_note
GROUP_FILLS = ("E2F2E7", "E2EBF7", "F8EED9", "F1E2F0")

MIN_SQUARE = 640        # квадратне відео — не менше 640×640
MIN_SIDE = 720          # горизонтальне чи вертикальне — коротша сторона від 720 (720p)
LOW_FILL = PatternFill("solid", fgColor="FFC7CE")
LOW_FONT = Font(bold=True, color="9C0006")


def format_resolution(resolution):
    """(1280, 720) → «1280x720»; невідомо (лише звук, файлу немає) — «»."""
    return f"{resolution[0]}x{resolution[1]}" if resolution else ""


def is_low_resolution(resolution):
    """Чи кадр нижчий за вимоги Rozetka до відео товару."""
    if not resolution:
        return False
    width, height = resolution
    return min(width, height) < (MIN_SQUARE if width == height else MIN_SIDE)


def status_label(state, text):
    """Коротко для фільтра в Excel: «Готово», «Копія», «Помилка»…"""
    text = text or ""
    if state == "done":
        if text.startswith(SAME_VIDEO):
            return "Те саме відео"
        if text.startswith("Готово — копія"):
            return "Копія"
        if text.startswith(ALREADY_NOTE):
            return "Уже було"
        if "без субтитрів" in text:
            return "Без субтитрів"
        return "Готово"
    return {"error": "Помилка", "cancelled": "Скасовано", "queued": "У черзі",
            "running": "Качається"}.get(state, state)


def _key(item):
    path = item.get("filepath") or ""
    return os.path.normcase(os.path.abspath(path)) if path and item.get("state") == "done" else None


def shared_groups(items):
    """Файл → {index, owner, others}: лише файли, що їх ділять кілька товарів.
    owner — ID, за яким названо файл (той, хто його качав); others — решта ID."""
    by_file = {}
    for item in items:
        key = _key(item)
        if key:
            by_file.setdefault(key, []).append(item)
    groups = {}
    for key, members in by_file.items():
        if len(members) < 2:
            continue
        owner = next((m for m in members if not (m.get("text") or "").startswith(SAME_VIDEO)),
                     members[0])
        groups[key] = {"index": len(groups), "owner": owner.get("product_id") or "",
                       "others": [m.get("product_id") or "" for m in members if m is not owner]}
    return groups


def details(item, groups):
    """«Деталі»: текст стану; у власника спільного файлу — для кого ще він."""
    text = item.get("text") or ""
    group = groups.get(_key(item))
    if group and not text.startswith(SAME_VIDEO):
        others = ", ".join(p for p in group["others"] if p)
        text += f"  ·  спільне відео ще для: {others}" if others else ""
    return text


def default_name(now=None):
    return time.strftime("zvit_%Y-%m-%d_%H%M.xlsx", time.localtime(now))


def write_report(path, items):
    """items — словники з ключами product_id, url, title, filepath, state,
    text, duration (секунди або None), resolution ((ширина, висота) або ()),
    ftp (шлях на FTP або стан заливання).
    Повертає path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Звіт"
    ws.append(HEADERS)
    groups = shared_groups(items)
    total, counted, low = 0, set(), 0
    for n, item in enumerate(items, 1):
        path_ = item.get("filepath") or ""
        shares = (item.get("text") or "").startswith(SAME_VIDEO)
        size = None
        if item.get("state") == "done" and path_ and os.path.isfile(path_) and not shares:
            size = round(os.path.getsize(path_) / 1024 / 1024, 1)
        duration = item.get("duration")
        resolution = tuple(item.get("resolution") or ()) if item.get("state") == "done" else ()
        # «Разом» — кожен файл один раз: товари з тим самим відео часу не додають.
        key = _key(item) or id(item)
        if duration and item.get("state") == "done" and key not in counted:
            counted.add(key)
            total += whole_seconds(duration)
        ws.append([
            n,
            item.get("product_id") or "",
            item.get("url") or "",
            item.get("title") or "",
            os.path.basename(path_),
            status_label(item.get("state"), item.get("text")),
            details(item, groups),
            format_min_sec(duration) if duration else "",
            format_resolution(resolution),
            size,
            item.get("ftp") or "",
        ])
        link = ws.cell(ws.max_row, len(HEADERS))
        if str(link.value or "").startswith("http"):
            link.hyperlink = link.value
            link.style = "Hyperlink"
        group = groups.get(_key(item))
        if group is not None:
            fill = PatternFill("solid", fgColor=GROUP_FILLS[group["index"] % len(GROUP_FILLS)])
            for cell in ws[ws.max_row]:
                cell.fill = fill
        if is_low_resolution(resolution):
            low += 1
            cell = ws.cell(ws.max_row, RESOLUTION_COL)
            cell.fill, cell.font = LOW_FILL, LOW_FONT
    if total or low:
        ws.append([])
    if total:
        ws.append(["", "", "", "Разом завантажено", "", "", "", format_min_sec(total), None])
        ws.cell(ws.max_row, 4).font = Font(bold=True)
        ws.cell(ws.max_row, 8).font = Font(bold=True)
    if low:
        ws.append(["", "", "", f"Нижче {MIN_SQUARE}x{MIN_SQUARE} / {MIN_SIDE}p", "", "", "",
                   "", low])
        for col in (4, RESOLUTION_COL):
            ws.cell(ws.max_row, col).font = LOW_FONT
        ws.cell(ws.max_row, RESOLUTION_COL).fill = LOW_FILL

    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(items) + 1}"
    for i in range(1, len(HEADERS) + 1):
        letter = get_column_letter(i)
        width = max(len(str(c.value)) for c in ws[letter] if c.value is not None)
        ws.column_dimensions[letter].width = min(MAX_WIDTH, width + 2)
    wb.save(path)
    return path
