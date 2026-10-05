"""Список «ID товару — посилання» з xlsx / csv: колонки знаходяться самі."""

import os
import tempfile
import unittest

from openpyxl import Workbook

from rtube import sheets

A = "https://www.youtube.com/watch?v=pn6mZ0Bcugo"
B = "https://www.youtube.com/watch?v=1G01ROKhAAw"
C = "https://www.youtube.com/watch?v=ZNMNKI4xPbY"


class SheetsTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def test_header_picks_id_column_over_other_numbers(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["Код продавця", "ID товару", "Назва", "Відео"])
        ws.append([123456, 590312170, "Килимок", "https://youtube.com/shorts/pn6mZ0Bcugo?si=x"])
        ws.append([123456, 580250272.0, "Мишка", "https://youtu.be/1G01ROKhAAw"])   # число як float
        ws.append([123456, None, "Без ID", "https://youtu.be/ZNMNKI4xPbY"])
        ws.append(["", "", "рядок без посилання", ""])
        # Гіперпосилання під текстом «відео»
        ws.append([123456, "593505175", "Гіперпосилання", "відео"])
        ws.cell(row=6, column=4).hyperlink = "https://www.youtube.com/watch?v=1zElSYng0Xg"
        wb.save(self.path("tovary.xlsx"))

        result = sheets.read_pairs(self.path("tovary.xlsx"))
        self.assertEqual(result.pairs, [("590312170", A), ("580250272", B), (None, C),
                                        ("593505175", "https://www.youtube.com/watch?v=1zElSYng0Xg")])
        self.assertEqual((result.with_ids, result.skipped), (3, 1))

    def test_no_header_url_first(self):
        wb = Workbook()
        wb.active.append([A, 590312170])
        wb.active.append([B, "580250272"])
        wb.save(self.path("a.xlsx"))
        self.assertEqual(sheets.read_pairs(self.path("a.xlsx")).pairs,
                         [("590312170", A), ("580250272", B)])

    def test_links_on_second_sheet(self):
        wb = Workbook()
        wb.active.title = "Зведення"
        wb.active.append(["нічого корисного"])
        ws = wb.create_sheet("Відео")
        ws.append(["ID", "URL"])
        ws.append([590312170, A])
        wb.save(self.path("b.xlsx"))
        result = sheets.read_pairs(self.path("b.xlsx"))
        self.assertEqual((result.pairs, result.sheet), ([("590312170", A)], "Відео"))

    def test_csv_semicolon_cp1251(self):
        with open(self.path("list.csv"), "w", encoding="cp1251", newline="") as f:
            f.write("ID товару;Посилання\r\n590312170;https://youtu.be/pn6mZ0Bcugo\r\n"
                    "580250272;https://youtu.be/1G01ROKhAAw\r\n")
        self.assertEqual(sheets.read_pairs(self.path("list.csv")).pairs,
                         [("590312170", A), ("580250272", B)])

    def test_txt_like_pasted_list(self):
        with open(self.path("list.txt"), "w", encoding="utf-8") as f:
            f.write("590312170 https://youtube.com/shorts/pn6mZ0Bcugo?si=abc ;\n")
        self.assertEqual(sheets.read_pairs(self.path("list.txt")).pairs, [("590312170", A)])

    def test_rozetka_template_takes_rozetka_code_not_price_id(self):
        # Як у справжньому .xls «Добавление видеообзора»: і код Rozetka, і ID
        # з прайс-листа продавця — обидва 9 цифр; плюс нерозривний пробіл у коді.
        rows = [["Код товару на ROZETKA", "Посилання на товар на сайті ROZETKA",
                 "ID товару у вашому прайс-листі", "Назва товару", "Посилання на відео"],
                ["610963253", "https://rozetka.com.ua/610963253/p610963253", "141903775",
                 "Замок-блокіратор", "https://www.youtube.com/shorts/pn6mZ0Bcugo"],
                ["610963244", "https://rozetka.com.ua/ua/610963244/p610963244/", "141890721",
                 "Блокіратор", "https://youtube.com/shorts/1G01ROKhAAw?si=K8Pie"]]
        rows[2][0] = sheets._cell_text("610963244\xa0")
        result = sheets.pairs_from_rows(rows)
        self.assertEqual(result.pairs, [("610963253", A), ("610963244", B)])

    def test_xls_supported_extension(self):
        self.assertIn(".xls", sheets.SUPPORTED)

    def test_errors(self):
        with self.assertRaises(ValueError):
            sheets.read_pairs(self.path("x.docx"))
        wb = Workbook()
        wb.active.append(["ID", "Назва"])
        wb.save(self.path("empty.xlsx"))
        with self.assertRaises(ValueError):
            sheets.read_pairs(self.path("empty.xlsx"))


if __name__ == "__main__":
    unittest.main()
